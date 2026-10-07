"""Pruebas de integridad de los archivos exportados."""
import csv
import importlib.util
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from planificador_procesos.core import Process, simulate
from planificador_procesos.exports import ExportData, export_csv, export_excel, export_pdf


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name)
        processes = (Process("Á", 0, 6), Process("B", 1, 3), Process("C", 2, 1))
        result = simulate(processes, "Round Robin", quantum=2, core_count=2)
        self.data = ExportData(processes, "Round Robin", 2, 2, result)

    def test_csv_separates_metadata_headers_and_numeric_results(self):
        for results in (False, True):
            with self.subTest(results=results):
                path = self.path / f"{results}.csv"
                export_csv(path, self.data, results=results)
                self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    rows = list(csv.reader(stream, delimiter=";"))
                self.assertIn("Round Robin", rows[0][0])
                self.assertIn("quantum = 2", rows[0][0])
                self.assertEqual(rows[2][0], "Proceso")
                self.assertEqual(rows[3][:3], ["Á", "0", "6"])
                self.assertEqual(len(rows[3]), 6 if results else 3)
                if results:
                    self.assertEqual([int(row[3]) for row in rows[3:]], [7, 4, 3])

    def test_csv_quotes_separator_newlines_and_formula_like_identifiers(self):
        processes = (Process("A;\nB", 0, 1), Process("=1+1", 0, 1))
        path = self.path / "text.csv"
        export_csv(path, ExportData(processes, "FCFS", None, 1))
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=";"))
        self.assertEqual(rows[3][0], "A;\nB")
        self.assertEqual(rows[4][0], "'=1+1")
        self.assertNotIn("quantum", rows[0][0])

    def test_failed_write_preserves_existing_file_and_cleans_temporary(self):
        path = self.path / "existing.csv"
        path.write_text("original", encoding="utf-8")
        with patch("planificador_procesos.fileio.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                export_csv(path, self.data)
        self.assertEqual(path.read_text(encoding="utf-8"), "original")
        self.assertEqual(list(self.path.iterdir()), [path])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl no instalado")
    def test_excel_has_one_sheet_numbers_metrics_and_all_cpu_assignments(self):
        from openpyxl import load_workbook
        path = self.path / "simulation.xlsx"
        export_excel(path, self.data)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Simulación"])
        sheet = workbook.active
        self.assertIn("quantum = 2", sheet["A1"].value)
        rows = list(sheet.values)
        result_index = next(i for i, row in enumerate(rows) if row[0] == "Resultados")
        self.assertEqual(list(rows[result_index + 2][:6]), ["Á", 0, 6, 7, 1, 7])
        average_index = next(i for i, row in enumerate(rows) if row[0] == "Espera media")
        self.assertAlmostEqual(rows[average_index][1], self.data.result.average_waiting_time)
        for core, assignment in enumerate(self.data.result.core_timeline, 1):
            row = next(row for row in rows if row[0] == f"CPU {core}")
            self.assertEqual(list(row[1:8]), [v if v is not None else "—" for v in assignment])
        self.assertEqual(sheet.page_setup.fitToHeight, 1)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl no instalado")
    def test_excel_process_ids_are_text_even_when_they_look_like_formulas(self):
        from openpyxl import load_workbook
        processes = (Process("=1+1", 0, 1), Process("001", 0, 1))
        data = ExportData(processes, "FCFS", None, 1, simulate(processes, "FCFS"))
        path = self.path / "identifiers.xlsx"
        export_excel(path, data)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        ids = [cell for row in workbook.active for cell in row if cell.value in ("=1+1", "001")]
        self.assertEqual(len(ids), 8)  # Tres tablas y dos asignaciones de CPU.
        self.assertTrue(all(cell.data_type == "s" for cell in ids))

    @unittest.skipUnless(importlib.util.find_spec("reportlab"), "reportlab no instalado")
    def test_pdf_is_exactly_one_page_even_for_long_simulations(self):
        for duration in (6, 100):
            processes = (Process("A", 0, duration), Process("B", 1, 3))
            data = ExportData(processes, "FCFS", None, 2, simulate(processes, "FCFS", core_count=2))
            path = self.path / f"{duration}.pdf"
            export_pdf(path, data)
            raw = path.read_bytes()
            self.assertTrue(raw.startswith(b"%PDF-"))
            self.assertEqual(len(re.findall(rb"/Type\s*/Page\b", raw)), 1)

    def test_results_cannot_be_exported_before_calculation(self):
        with self.assertRaises(ValueError):
            export_csv(self.path / "missing.csv", ExportData(self.data.processes, "FCFS", None, 1), results=True)
        self.assertFalse((self.path / "missing.csv").exists())


if __name__ == "__main__":
    unittest.main()
