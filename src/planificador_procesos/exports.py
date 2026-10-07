"""Exportaciones de datos e informes, independientes de Tkinter.

PDF y XLSX importan sus dependencias bajo demanda. Los archivos se preparan
junto al destino y se reemplazan sólo después de generarlos correctamente.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from .core import Process, SimulationResult
from .fileio import atomic_destination as _atomic_destination
from .metadata import APP_NAME, AUTHOR, REPOSITORY_URL

PROCESS_HEADERS = ("Proceso", "Llegada", "Duración")
RESULT_HEADERS = (*PROCESS_HEADERS, "Finalización", "Espera", "Retorno")


@dataclass(frozen=True)
class ExportData:
    processes: tuple[Process, ...]
    algorithm: str
    quantum: int | None
    core_count: int
    result: SimulationResult | None = None

    @property
    def context(self) -> str:
        policy = self.algorithm
        if policy == "Round Robin":
            policy += f" · quantum = {self.quantum}"
        return f"{policy} · núcleos = {self.core_count}"

    def table(self, results: bool = False) -> list[list[str | int]]:
        if results:
            if self.result is None:
                raise ValueError("Calcule la planificación antes de exportar resultados.")
            return [list(RESULT_HEADERS)] + [
                [p.process_id, p.arrival_time, p.duration_time, p.completed_time,
                 p.waiting_time, p.turnaround_time] for p in self.result.processes
            ]
        return [list(PROCESS_HEADERS)] + [
            [p.process_id, p.arrival_time, p.duration_time] for p in self.processes
        ]

    def timelines(self) -> tuple[list[list], list[list]]:
        if self.result is None:
            raise ValueError("Calcule la planificación antes de exportar el informe.")
        cycles = ["Proceso", *range(1, self.result.makespan + 1)]
        processes = [cycles] + [[name, *cells] for name, cells in self.result.timeline.items()]
        cores = [["Núcleo", *cycles[1:]]] + [
            [f"CPU {i}", *(p if p is not None else "—" for p in cells)]
            for i, cells in enumerate(self.result.core_timeline, 1)
        ]
        return processes, cores


def _csv_value(value):
    # Los identificadores son texto, nunca fórmulas al abrirlos con Excel.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_csv(path: str | Path, data: ExportData, *, results: bool = False) -> None:
    rows = data.table(results)
    with _atomic_destination(path) as temporary:
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow([data.context])
            writer.writerow([])
            writer.writerows([_csv_value(value) for value in row] for row in rows)


def export_excel(path: str | Path, data: ExportData) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    timeline, cores = data.timelines()
    if len(timeline[0]) > 16384 or 2 * len(data.processes) + len(timeline) + len(cores) + 24 > 1048576:
        raise ValueError("La simulación supera los límites de filas o columnas de Excel.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Simulación"
    sheet.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D5DEE7")
    row_index = 1
    max_column = max(6, len(timeline[0]))

    def put(row, column, value):
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            cell.data_type = "s"
        cell.font = Font(name="Calibri", size=11, color="213547")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        return cell

    def heading(text, *, title=False):
        nonlocal row_index
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=6)
        cell = put(row_index, 1, text)
        cell.font = Font(name="Calibri", size=14 if title else 12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24445C")
        sheet.row_dimensions[row_index].height = 32
        row_index += 1

    def table(rows, *, kind="table"):
        nonlocal row_index
        arrivals = {p.process_id: p.arrival_time for p in data.processes}
        for index, row in enumerate(rows):
            for column, value in enumerate(row, 1):
                cell = put(row_index, column, value)
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(horizontal="left" if column == 1 else "center",
                                           vertical="center", wrap_text=True)
                color = "F2F6FA" if index % 2 == 0 else "FFFFFF"
                if index == 0:
                    color = "DBE7F0"
                    cell.font = Font(name="Calibri", size=11, bold=True, color="213547")
                elif kind == "cores" and column > 1 and value != "—":
                    color = "DCEEFB"
                elif kind == "timeline":
                    if (column == 1 and arrivals[row[0]] == 0) or (column > 1 and arrivals[row[0]] == column - 1):
                        color = "D9D9D9"
                    elif value == "X":
                        color = "DDF2E0"
                    elif value == "O":
                        color = "FFF0C2"
                cell.fill = PatternFill("solid", fgColor=color)
            longest = max(len(str(value)) for value in row)
            sheet.row_dimensions[row_index].height = max(26, 16 * ((longest + 15) // 16))
            row_index += 1
        row_index += 1

    heading(data.context, title=True)  # A1 conserva política, quantum y núcleos.
    put(row_index, 1, APP_NAME)
    row_index += 2
    heading("Procesos")
    table(data.table())
    heading("Resultados")
    table(data.table(True))
    put(row_index, 1, "Espera media")
    put(row_index, 2, data.result.average_waiting_time).number_format = "0.00"
    put(row_index, 4, "Retorno medio")
    put(row_index, 5, data.result.average_turnaround_time).number_format = "0.00"
    row_index += 2
    heading("Tabla de ciclos")
    put(row_index, 1, "X: ejecuta · O: espera · Gris: llegada")
    sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=6)
    row_index += 1
    table(timeline, kind="timeline")
    heading("Diagrama de núcleos")
    table(cores, kind="cores")
    put(row_index, 1, REPOSITORY_URL)
    sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=6)
    sheet.column_dimensions["A"].width = 28
    for column in range(2, max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 17 if column <= 6 else 9
    sheet.freeze_panes = "B2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.3, bottom=.3)
    sheet.print_area = f"A1:{get_column_letter(max_column)}{row_index}"
    with _atomic_destination(path) as temporary:
        workbook.save(temporary)
    workbook.close()


def _pdf_layout(data: ExportData):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    timeline, cores = data.timelines()
    label_width = min(180, max(90, *(stringWidth(p.process_id, "Helvetica", 9) + 14 for p in data.processes)))
    cycle_width = max(24, min(100, max(stringWidth(str(v), "Helvetica", 8) + 10 for row in cores for v in row[1:])))
    width = max(770, label_width + data.result.makespan * cycle_width)
    normal = ParagraphStyle("normal", fontName="Helvetica", fontSize=9, leading=12)
    cell_style = ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10)
    title = ParagraphStyle("title", parent=normal, fontSize=18, leading=22, textColor=colors.HexColor("#24445c"))
    section = ParagraphStyle("section", parent=normal, fontSize=11, leading=14, spaceBefore=4)
    flows = [Paragraph(APP_NAME, title), Spacer(1, 6), Paragraph(escape(data.context), normal), Spacer(1, 12)]

    def add_table(label, rows, widths, *, kind="table"):
        flows.extend([Paragraph(label, section), Spacer(1, 4)])
        wrapped = [[Paragraph(escape(str(value)), cell_style) for value in row] for row in rows]
        table = Table(wrapped, colWidths=widths, hAlign="LEFT")
        commands = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe7f0")),
                    ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d5dee7")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
        if kind == "table":
            commands.append(("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fa")]))
        else:
            arrivals = {p.process_id: p.arrival_time for p in data.processes}
            for row_index, row in enumerate(rows[1:], 1):
                for column, value in enumerate(row):
                    color = None
                    if kind == "cores" and column and value != "—":
                        color = "#dceefb"
                    elif kind == "timeline":
                        if (column == 0 and arrivals[row[0]] == 0) or (column > 0 and arrivals[row[0]] == column):
                            color = "#d9d9d9"
                        elif column and value in ("X", "O"):
                            color = "#ddf2e0" if value == "X" else "#fff0c2"
                    if color:
                        commands.append(("BACKGROUND", (column, row_index), (column, row_index), colors.HexColor(color)))
        table.setStyle(TableStyle(commands))
        flows.extend([table, Spacer(1, 10)])

    add_table("Procesos", data.table(), [max(160, label_width), 100, 100])
    add_table("Resultados", data.table(True), [max(160, label_width), 100, 100, 110, 100, 100])
    flows.extend([Paragraph(f"Espera media: {data.result.average_waiting_time:.2f} · "
                            f"Retorno medio: {data.result.average_turnaround_time:.2f} · "
                            f"Duración total: {data.result.makespan} ciclos", normal), Spacer(1, 10)])
    add_table("Tabla de ciclos · X: ejecuta · O: espera · Gris: llegada", timeline,
              [label_width] + [cycle_width] * data.result.makespan, kind="timeline")
    add_table("Diagrama de núcleos", cores, [label_width] + [cycle_width] * data.result.makespan, kind="cores")
    dimensions = [flow.wrap(width, 10000000) for flow in flows]
    height = sum(h for _, h in dimensions)
    return flows, width, height, dimensions


def pdf_minimum_font_size(data: ExportData) -> float:
    from reportlab.lib.pagesizes import A4, landscape
    _, width, height, _ = _pdf_layout(data)
    page_width, page_height = landscape(A4)
    return 8 * min(1, (page_width - 56) / width, (page_height - 80) / height)


def export_pdf(path: str | Path, data: ExportData) -> None:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen.canvas import Canvas

    flows, width, height, dimensions = _pdf_layout(data)
    page_width, page_height = landscape(A4)
    scale = min(1, (page_width - 56) / width, (page_height - 80) / height)
    with _atomic_destination(path) as temporary:
        canvas = Canvas(str(temporary), pagesize=(page_width, page_height))
        canvas.setTitle(f"{APP_NAME} · {data.context}")
        canvas.setAuthor(AUTHOR)
        canvas.saveState()
        canvas.translate(28, page_height - 28)
        canvas.scale(scale, scale)
        y = 0
        for flow, (_, flow_height) in zip(flows, dimensions):
            y -= flow_height
            flow.drawOn(canvas, 0, y)
        canvas.restoreState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColorRGB(.3, .3, .3)
        canvas.drawString(28, 20, f"{AUTHOR} · {REPOSITORY_URL}")
        canvas.showPage()
        canvas.save()
