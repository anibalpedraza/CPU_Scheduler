"""Comprobación del paquete real, sin consola ni diálogos interactivos."""
from __future__ import annotations

import csv
import json
from pathlib import Path
import re
import platform
import sys
import tempfile
import traceback

from .metadata import VERSION


def run_check(report_path: Path) -> int:
    report = {"version": VERSION, "frozen": bool(getattr(sys, "frozen", False)),
              "success": False, "architecture": platform.machine().lower(), "checks": []}
    root = None
    try:
        import tkinter as tk
        from openpyxl import load_workbook
        from .app import SchedulerApp
        from .core import SUPPORTED_ALGORITHMS
        from .exports import ExportData, export_csv, export_excel, export_pdf
        from .gui_checks import check_dialog_lifecycle, check_timeline_contrast
        from .persistence import load_simulation, save_simulation
        from tkinter import messagebox

        def fail_dialog(title, message, **_kwargs):
            raise RuntimeError(f"{title}: {message}")

        messagebox.showerror = fail_dialog
        root = tk.Tk()
        root.withdraw()
        # Reproduce el texto blanco que puede imponer un tema oscuro incluso en Windows.
        root.option_add("*Label.foreground", "#ffffff")
        root.option_add("*Label.background", "#303030")
        root.report_callback_exception = lambda _kind, value, _tb: fail_dialog("Tkinter", str(value))
        app = SchedulerApp(root)
        root.update()
        labels = [app.menu_bar.entrycget(i, "label")
                  for i in range(app.menu_bar.index("end") + 1)
                  if app.menu_bar.type(i) == "cascade"]
        if labels != ["Archivo", "Edición", "Ayuda"]:
            raise RuntimeError(f"Menús inesperados: {labels}")
        from .licensing import license_text
        notices = license_text()
        if "MIT License" not in notices:
            raise RuntimeError("Falta la licencia de la aplicación")
        if report["frozen"] and any(name not in notices.lower()
                                    for name in ("pyinstaller", "reportlab", "openpyxl", "pillow")):
            raise RuntimeError("Faltan avisos de componentes de terceros")
        root.deiconify()
        root.update()
        app.show_licenses()
        root.update()
        dialog = app._licenses_dialog
        app.show_licenses()
        if not dialog.winfo_ismapped() or app._licenses_dialog is not dialog or root.grab_current() is not None:
            raise RuntimeError("La ventana de licencias se duplica o bloquea la aplicación")
        dialog.destroy()
        root.update()
        root.withdraw()
        report["checks"].append("Licencias integradas y consultables")
        report["checks"].append("Tkinter, Tcl/Tk e interfaz con menús")
        app.load_example()
        for cores in (1, 2):
            app.core_count_var.set(str(cores))
            for algorithm in SUPPORTED_ALGORITHMS:
                app.algorithm_var.set(algorithm)
                app.calculate()
                root.update()
                if app.last_result is None or len(app.result_table.get_children()) != 3:
                    raise RuntimeError(f"Falló {algorithm}, {cores} núcleos")
                report["checks"].append(f"{algorithm}: {cores} núcleos")
                check_timeline_contrast(app)
        report["checks"].append("Contraste de ciclos y CPU con texto blanco por defecto")
        check_dialog_lifecycle(app)
        report["checks"].append("Acerca de y CSV: apertura, eventos, cierre y reapertura sin bloqueo")
        with tempfile.TemporaryDirectory(prefix="cpu-scheduler-check-") as directory:
            target = Path(directory)
            document = app._simulation_document()
            save_simulation(target / "simulación.json", document)
            if load_simulation(target / "simulación.json") != document:
                raise RuntimeError("El documento JSON no se recupera correctamente")
            report["checks"].append("Guardar y recuperar JSON")
            data = ExportData(document.processes, document.algorithm, document.quantum,
                              document.core_count, app.last_result)
            export_csv(target / "resultados.csv", data, results=True)
            with (target / "resultados.csv").open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream, delimiter=";"))
            if len(rows) != 6 or rows[2][0] != "Proceso" or rows[3][0] != "A":
                raise RuntimeError("Contenido CSV incorrecto")
            report["checks"].append("Exportación CSV")
            export_pdf(target / "informe.pdf", data)
            raw = (target / "informe.pdf").read_bytes()
            if not raw.startswith(b"%PDF-") or len(re.findall(rb"/Type\s*/Page\b", raw)) != 1:
                raise RuntimeError("El PDF no contiene una página válida")
            report["checks"].append("Exportación PDF de una página")
            export_excel(target / "informe.xlsx", data)
            workbook = load_workbook(target / "informe.xlsx")
            try:
                if workbook.sheetnames != ["Simulación"] or "núcleos = 2" not in workbook.active["A1"].value:
                    raise RuntimeError("Contenido Excel incorrecto")
            finally:
                workbook.close()
            report["checks"].append("Exportación y lectura Excel")
        report["success"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    finally:
        if root is not None:
            root.destroy()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if report["success"] else 1
